from typing import List, Dict
import os
from pathlib import Path
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from app.core.config import settings

class DocumentProcessor:
    """Process different document types and extract text with page numbers"""
    
    @staticmethod
    def process_pdf(file_path: str) -> List[Dict[str, any]]:
        """Extract text from PDF with page numbers"""
        chunks = []
        
        try:
            pdf_reader = PdfReader(file_path)
            
            for page_num, page in enumerate(pdf_reader.pages, start=1):
                text = page.extract_text()
                
                if text and text.strip():
                    # Split into smaller chunks while preserving page info
                    text_chunks = DocumentProcessor._chunk_text(text)
                    
                    for chunk_text in text_chunks:
                        chunks.append({
                            "text": chunk_text,
                            "page": page_num,
                            "type": "pdf"
                        })
        except Exception as e:
            raise ValueError(f"Error processing PDF: {str(e)}")
        
        return chunks
    
    @staticmethod
    def process_docx(file_path: str) -> List[Dict[str, any]]:
        """Extract text from DOCX with page approximation"""
        chunks = []
        
        try:
            doc = Document(file_path)
            
            # Approximate page breaks (roughly 500 words per page)
            words_per_page = 500
            current_word_count = 0
            current_page = 1
            current_text = []
            
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                
                word_count = len(text.split())
                current_word_count += word_count
                current_text.append(text)
                
                # Approximate page break
                if current_word_count >= words_per_page:
                    full_text = " ".join(current_text)
                    text_chunks = DocumentProcessor._chunk_text(full_text)
                    
                    for chunk_text in text_chunks:
                        chunks.append({
                            "text": chunk_text,
                            "page": current_page,
                            "type": "docx"
                        })
                    
                    current_page += 1
                    current_word_count = 0
                    current_text = []
            
            # Process remaining text
            if current_text:
                full_text = " ".join(current_text)
                text_chunks = DocumentProcessor._chunk_text(full_text)
                
                for chunk_text in text_chunks:
                    chunks.append({
                        "text": chunk_text,
                        "page": current_page,
                        "type": "docx"
                    })
        except Exception as e:
            raise ValueError(f"Error processing DOCX: {str(e)}")
        
        return chunks
    
    @staticmethod
    def process_xlsx(file_path: str) -> List[Dict[str, any]]:
        """Extract text from XLSX with sheet information"""
        chunks = []
        
        try:
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows_text = []
                
                for row in sheet.iter_rows(values_only=True):
                    if row:
                        row_text = " | ".join([str(cell) for cell in row if cell is not None])
                        if row_text.strip():
                            rows_text.append(row_text)
                
                if rows_text:
                    full_text = "\n".join(rows_text)
                    text_chunks = DocumentProcessor._chunk_text(full_text)
                    
                    for chunk_text in text_chunks:
                        chunks.append({
                            "text": chunk_text,
                            "page": sheet_name,  # Use sheet name as "page"
                            "type": "xlsx"
                        })
            
            workbook.close()
        except Exception as e:
            raise ValueError(f"Error processing XLSX: {str(e)}")
        
        return chunks
    
    @staticmethod
    def _chunk_text(text: str) -> List[str]:
        """Split text into overlapping chunks"""
        if not text:
            return []
        
        words = text.split()
        if len(words) <= settings.CHUNK_SIZE:
            return [text]
        
        chunks = []
        
        for i in range(0, len(words), settings.CHUNK_SIZE - settings.CHUNK_OVERLAP):
            chunk = " ".join(words[i:i + settings.CHUNK_SIZE])
            if chunk.strip():
                chunks.append(chunk)
        
        return chunks if chunks else [text]
    
    @staticmethod
    def process_document(file_path: str) -> List[Dict[str, any]]:
        """Process document based on file extension"""
        ext = Path(file_path).suffix.lower()
        
        if ext == '.pdf':
            return DocumentProcessor.process_pdf(file_path)
        elif ext == '.docx':
            return DocumentProcessor.process_docx(file_path)
        elif ext == '.xlsx':
            return DocumentProcessor.process_xlsx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")